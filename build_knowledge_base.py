from pptx import Presentation
from pypdf import PdfReader
from sentence_transformers import SentenceTransformer
import faiss
import os
import pickle


# -----------------------------
# Settings
# -----------------------------

knowledge_folder = "knowledge_base"

documents = []


# -----------------------------
# Read PowerPoint files
# -----------------------------

for filename in os.listdir(knowledge_folder):

    if (
        filename.endswith(".pptx")
        and not filename.startswith("~$")
    ):

        filepath = os.path.join(
            knowledge_folder,
            filename
        )

        print(f"Reading PPT: {filename}")

        presentation = Presentation(filepath)

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1
        ):

            slide_text = []

            for shape in slide.shapes:

                if hasattr(shape, "text"):

                    text = shape.text.strip()

                    if text:
                        slide_text.append(text)

            if slide_text:

                documents.append({
                    "text": "\n".join(slide_text),
                    "source": filename,
                    "location": f"Slide {slide_number}"
                })


# -----------------------------
# Read PDF files
# -----------------------------

for filename in os.listdir(knowledge_folder):

    if filename.lower().endswith(".pdf"):

        filepath = os.path.join(
            knowledge_folder,
            filename
        )

        print(f"Reading PDF: {filename}")

        reader = PdfReader(filepath)

        for page_number, page in enumerate(
            reader.pages,
            start=1
        ):

            text = page.extract_text()

            if text and text.strip():

                documents.append({
                    "text": text.strip(),
                    "source": filename,
                    "location": f"Page {page_number}"
                })


# -----------------------------
# Check documents
# -----------------------------

print(
    f"\nTotal documents extracted: {len(documents)}"
)


if not documents:

    print(
        "No PDF or PPTX files found in knowledge_base."
    )

    exit()


# -----------------------------
# Create chunks
# -----------------------------

chunks = []

chunk_size = 150

for document in documents:

    words = document["text"].split()

    for i in range(
        0,
        len(words),
        chunk_size
    ):

        chunk_text = " ".join(
            words[i:i + chunk_size]
        )

        if chunk_text.strip():

            chunks.append({
                "text": chunk_text,
                "source": document["source"],
                "location": document["location"]
            })


print(
    f"Total chunks created: {len(chunks)}"
)


# -----------------------------
# Create embeddings
# -----------------------------

print(
    "\nLoading embedding model..."
)

model = SentenceTransformer(
    "all-MiniLM-L6-v2"
)

texts = [
    chunk["text"]
    for chunk in chunks
]

embeddings = model.encode(
    texts,
    convert_to_numpy=True,
    show_progress_bar=True
)


# -----------------------------
# Create FAISS index
# -----------------------------

dimension = embeddings.shape[1]

index = faiss.IndexFlatL2(
    dimension
)

index.add(embeddings)

print(
    f"\nFAISS index created with "
    f"{index.ntotal} vectors."
)


# -----------------------------
# Save FAISS index
# -----------------------------

faiss.write_index(
    index,
    "knowledge_base/faiss_index.bin"
)


# -----------------------------
# Save chunks and metadata
# -----------------------------

with open(
    "knowledge_base/chunks.pkl",
    "wb"
) as file:

    pickle.dump(
        chunks,
        file
    )


print(
    "\nKnowledge base created successfully!"
)

print(
    "\nCreated files:"
)

print(
    "knowledge_base/faiss_index.bin"
)

print(
    "knowledge_base/chunks.pkl"
)